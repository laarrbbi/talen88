"""Document text extraction + deterministic field suggestion (Document Import).

Real, AI-free parsing for the "add employee" flow over ANY employee document — a CV, an
employment contract, an offer letter, a payslip, or a performance review (PDF via pypdf, Word
via python-docx, plain text directly). The document is classified (or the caller names its
type) and the matching heuristics propose the fields that document carries:
  * CV → name, email, phone, role/title, location, skills, years of experience, prior employers;
  * contract / offer / payslip → base salary + currency, hire/start date, employment type, title;
  * review → performance rating, goal attainment.
Each suggestion carries a confidence. Nothing is fabricated — a field the heuristics cannot find
is simply omitted, and the operator confirms or edits every suggestion (and supplies what no
document carries — division, level) before anything is written. Several documents can be uploaded
for one person; the UI merges their suggestions.

Heuristics only: no model call, so the same document always yields the same suggestions.
"""
from __future__ import annotations

import io
import re
from typing import Any

from .ingestion import normalize

# Supported upload extensions -> how we read them.
SUPPORTED_EXTS = (".pdf", ".docx", ".txt", ".md", ".text")

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
_PHONE_RE = re.compile(r"(?<!\w)(\+?\d[\d\s().\-]{7,}\d)(?!\w)")
_YEARS_RE = re.compile(r"(\d{1,2})\+?\s*(?:years|yrs)\b", re.IGNORECASE)

# Role keywords used to spot a job-title line near the top of a CV (domain-agnostic).
_ROLE_WORDS = (
    "engineer", "developer", "manager", "analyst", "designer", "consultant",
    "director", "lead", "scientist", "architect", "specialist", "associate",
    "administrator", "officer", "accountant", "recruiter", "coordinator",
    "strategist", "trader", "advisor", "controller", "head", "vp", "president",
    "researcher", "marketer", "salesperson", "technician", "product manager",
)

# A modest, multi-domain skill lexicon for the fallback skills scan (when there is no
# explicit "Skills" section). Matched case-insensitively as whole tokens/phrases.
_SKILL_LEXICON = (
    "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust", "ruby",
    "sql", "nosql", "postgres", "mysql", "mongodb", "excel", "powerpoint", "word",
    "react", "angular", "vue", "node", "django", "flask", "spring", "kubernetes",
    "docker", "aws", "azure", "gcp", "terraform", "pandas", "numpy", "pytorch",
    "tensorflow", "tableau", "power bi", "salesforce", "sap", "figma", "sketch",
    "photoshop", "illustrator", "git", "linux", "machine learning", "data analysis",
    "project management", "agile", "scrum", "accounting", "financial modeling",
    "risk management", "compliance", "marketing", "seo", "copywriting", "recruiting",
)

_NAME_STOPWORDS = {"curriculum", "vitae", "resume", "résumé", "cv", "profile", "contact"}
# Words that mark a document title / header line — never a person's name.
_NAME_REJECT = _NAME_STOPWORDS | {
    "employment", "contract", "agreement", "offer", "letter", "payslip", "payroll",
    "review", "appraisal", "performance", "annual", "confidential", "employee",
    "salary", "compensation", "statement", "certificate", "reference", "position",
}
# Section headers that bound a standalone "Skills" list and must never be read as a name.
_SECTION_WORDS = {
    "education", "experience", "summary", "projects", "certifications", "languages",
    "interests", "references", "work", "employment", "achievements", "awards",
    "profile", "objective", "contact", "skills",
}
_EDU_LEVELS = (
    ("doctorate", ("ph.d", "phd", "doctor of philosophy", "doctorate", "dphil")),
    ("master", ("master", "msc", "m.sc", "m.s.", "mba", "m.eng", "ma ", "ms ")),
    ("bachelor", ("bachelor", "bsc", "b.sc", "b.s.", "b.a", "ba ", "bs ", "b.eng", "btech")),
    ("associate", ("associate degree", "a.a.", "associate's")),
    ("high_school", ("high school", "secondary school", "ged")),
)

# ---- document-type classification + money/date/rating extraction --------------------
DOC_TYPES = ("cv", "contract", "offer", "payslip", "review")

_FILENAME_HINTS = {
    "contract": ("contract", "cdi", "agreement"),
    "offer": ("offer",),
    "payslip": ("payslip", "payroll", "salary"),
    "review": ("review", "appraisal", "performance"),
    "cv": ("cv", "resume", "résumé", "curriculum"),
}
_CONTENT_KEYWORDS = {
    "contract": ("employment contract", "contract of employment", "this agreement",
                 "terms of employment", "permanent contract", "indefinite term"),
    "offer": ("offer of employment", "we are pleased to offer", "offer letter",
              "letter of offer", "job offer"),
    "payslip": ("payslip", "pay slip", "salary slip", "net pay", "gross pay", "payroll"),
    "review": ("performance review", "performance appraisal", "annual review",
               "goal attainment", "performance rating", "appraisal"),
    "cv": ("curriculum vitae", "résumé", "professional experience", "work experience"),
}

# Currency sign/code -> ISO code (matches data.canonical.CURRENCIES).
_CURRENCY_SIGNS = {"S$": "SGD", "HK$": "HKD", "$": "USD", "£": "GBP", "€": "EUR", "¥": "JPY"}
_SALARY_WORDS = ("salary", "base", "annual", "gross", "compensation", "remuneration", "pay")
# Optional currency sign/code, a grouped number, optional trailing 'k'.
_MONEY_RE = re.compile(
    r"(S\$|HK\$|[$£€¥])?\s*(USD|GBP|EUR|SGD|HKD|JPY)?\s*"
    r"([0-9][0-9,]{2,}(?:\.[0-9]{2})?)\s*(k)?\b", re.IGNORECASE)
_DATE_TOKEN = r"([0-9]{1,4}[/\-.][0-9]{1,2}[/\-.][0-9]{1,4}|[A-Za-z]{3,9}\s+[0-9]{1,2},?\s+[0-9]{4})"
_HIRE_LABELS = ("start date", "commencement date", "commencement", "hire date", "date of joining",
                "joining date", "effective date", "effective", "date of commencement")
_DATE_RANGE_RE = re.compile(r"\b(?:19|20)\d{2}\s*[-–—to]+\s*(?:(?:19|20)\d{2}|present|current)\b",
                            re.IGNORECASE)

# Spoken-language lexicon for the fallback scan (the explicit "Languages:" section wins).
_LANGUAGES = (
    "english", "french", "spanish", "german", "mandarin", "chinese", "cantonese", "arabic",
    "portuguese", "italian", "dutch", "russian", "japanese", "korean", "hindi", "bengali",
    "punjabi", "turkish", "polish", "swedish", "norwegian", "danish", "finnish", "greek",
    "hebrew", "thai", "vietnamese", "indonesian", "malay", "tagalog", "urdu", "farsi", "persian",
)
# Degree subject -> canonical education_field (checked in this order; first hit wins).
_EDU_FIELD_KEYWORDS = (
    ("marketing", ("marketing", "advertising", "communications")),
    ("technical_degree", ("computer science", "software", "engineering", "information technology",
                          "data science", "mathematics", "physics")),
    ("life_sciences", ("biology", "chemistry", "life sciences", "biochemistry", "biotechnology")),
    ("medical", ("medicine", "medical", "nursing", "pharmacy", "public health")),
    ("human_resources", ("human resources", "organizational psychology")),
    ("business", ("business", "mba", "finance", "economics", "accounting", "management")),
)
_UNI_RE = re.compile(
    r"\b([A-Z][A-Za-z.&'-]+(?:\s+[A-Z][A-Za-z.&'-]+){0,4}\s+"
    r"(?:University|College|Institute of Technology|Polytechnic))\b")
_UNI_OF_RE = re.compile(r"\b((?:University|College|Institute)\s+of\s+"
                        r"[A-Z][A-Za-z'-]+(?:\s+[A-Z][A-Za-z'-]+){0,3})\b")


class UnsupportedDocument(ValueError):
    """Raised when an upload is not a CV format we can read."""


def extract_text(filename: str, raw: bytes) -> str:
    """Extract plain text from an uploaded CV. Dispatch on the filename extension.

    Raises UnsupportedDocument for an extension we cannot read. PDF/DOCX extraction is
    defensive: a corrupt or image-only document yields an empty string rather than raising.
    """
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _pdf_text(raw)
    if name.endswith(".docx"):
        return _docx_text(raw)
    if name.endswith((".txt", ".md", ".text")):
        return raw.decode("utf-8", errors="ignore")
    raise UnsupportedDocument(
        f"unsupported document type; upload one of {', '.join(SUPPORTED_EXTS)}"
    )


def _pdf_text(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is pinned in requirements
        raise UnsupportedDocument("PDF support is unavailable (pypdf not installed)") from exc
    try:
        reader = PdfReader(io.BytesIO(raw))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        return ""  # encrypted / image-only / malformed — caller falls back to manual entry


def _docx_text(raw: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is pinned in requirements
        raise UnsupportedDocument("Word support is unavailable (python-docx not installed)") from exc
    try:
        document = docx.Document(io.BytesIO(raw))
        return "\n".join(p.text for p in document.paragraphs)
    except Exception:
        return ""


# --- field heuristics ---------------------------------------------------------------

def _looks_like_name(line: str) -> bool:
    if "@" in line or any(ch.isdigit() for ch in line):
        return False
    words = line.split()
    if not (2 <= len(words) <= 4):
        return False
    if any(w.lower().strip(".,:") in _NAME_REJECT for w in words):
        return False
    for w in words:
        core = w.replace("-", "").replace("'", "").replace(".", "")  # tolerate middle initials
        if not core or not core[0].isupper() or not core.isalpha():
            return False
    return True


def _guess_name(lines: list[str]) -> str | None:
    # A labeled name wins ("Employee: Dana Lee", "Name: …"), common in contracts/reviews.
    for line in lines[:14]:
        m = re.match(r"(?:employee|name|candidate|full\s+name)\s*[:\-]\s*(.+)", line, re.IGNORECASE)
        if m:
            cand = m.group(1).strip()
            if _looks_like_name(cand):
                return cand.title() if cand.isupper() else cand
    # Otherwise the first top line that has a name's shape and no role/title word.
    for line in lines[:8]:
        if any(rw in line.lower() for rw in _ROLE_WORDS):
            continue
        if _looks_like_name(line):
            return line.title() if line.isupper() else line
    return None


def _guess_role(lines: list[str]) -> str | None:
    # Explicit labels win.
    for line in lines[:25]:
        m = re.match(r"(?:current\s+)?(?:role|title|position)\s*[:\-]\s*(.+)", line, re.IGNORECASE)
        if m and m.group(1).strip():
            return m.group(1).strip()
    # Otherwise the first short line (near the top) that reads like a job title.
    for line in lines[:12]:
        low = line.lower()
        if 2 <= len(line.split()) <= 6 and any(w in low for w in _ROLE_WORDS):
            return line.strip(" .-")
    return None


def _guess_location(lines: list[str]) -> str | None:
    for line in lines[:30]:
        m = re.match(r"(?:location|based in|address)\s*[:\-]\s*(.+)", line, re.IGNORECASE)
        if m and m.group(1).strip():
            return m.group(1).strip()
    return None


def _guess_skills(text: str, lines: list[str]) -> list[str]:
    # Prefer an explicit "Skills" section.
    for i, line in enumerate(lines):
        m = re.match(r"\s*(?:technical\s+|core\s+)?skills\s*[:\-]?\s*(.*)$", line, re.IGNORECASE)
        if not m:
            continue
        inline = m.group(1).strip()
        if inline:
            blob = inline                          # "Skills: a, b, c" — list is on this line
        else:
            collected: list[str] = []              # standalone "Skills:" header — read following
            for nxt in lines[i + 1:i + 5]:         # lines until the next section header
                first = nxt.split()[0].lower().strip(":,") if nxt.split() else ""
                if first in _SECTION_WORDS:
                    break
                collected.append(nxt)
            blob = " , ".join(collected)
        parts = re.split(r"[,;|•·]+", blob)
        skills = _dedupe_keep_order([p.strip() for p in parts if 1 < len(p.strip()) <= 30])
        if skills:
            return skills[:12]
    # Fallback: scan for known skills anywhere in the text.
    low = text.lower()
    found = [s for s in _SKILL_LEXICON if re.search(rf"(?<!\w){re.escape(s)}(?!\w)", low)]
    return _dedupe_keep_order([s.title() if s.islower() else s for s in found])[:12]


def _guess_years(text: str) -> int | None:
    matches = [int(m) for m in _YEARS_RE.findall(text)]
    plausible = [m for m in matches if 0 < m <= 50]
    return max(plausible) if plausible else None


def _guess_education(text: str) -> str | None:
    low = text.lower()
    for level, needles in _EDU_LEVELS:
        if any(n in low for n in needles):
            return level
    return None


def _dedupe_keep_order(items: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for it in items:
        key = it.lower()
        if key not in seen:
            seen.add(key)
            out.append(it)
    return out


def detect_doc_type(filename: str, text: str) -> str:
    """Classify a document from its filename then its content. Falls back to 'cv'."""
    fn = (filename or "").lower()
    for dtype, hints in _FILENAME_HINTS.items():
        if any(h in fn for h in hints):
            return dtype
    low = (text or "").lower()
    best, score = "cv", 0
    for dtype, kws in _CONTENT_KEYWORDS.items():
        c = sum(1 for kw in kws if kw in low)
        if c > score:
            best, score = dtype, c
    return best


def _guess_money(text: str) -> tuple[float, str] | None:
    """Largest salary-like amount + its currency. A bare number counts only with a salary word
    nearby (so we never read a year or a phone number as pay)."""
    low = text.lower()
    best: tuple[float, str] | None = None
    for m in _MONEY_RE.finditer(text):
        sign, code, num, k = m.group(1), m.group(2), m.group(3), m.group(4)
        if not sign and not code and not k:
            ctx = low[max(0, m.start() - 40): m.end() + 12]
            if not any(w in ctx for w in _SALARY_WORDS):
                continue
        try:
            amount = float(num.replace(",", ""))
        except ValueError:
            continue
        if k:
            amount *= 1000
        if amount < 1000:                      # too small to be annual pay
            continue
        ccy = _CURRENCY_SIGNS.get(sign) if sign else (code.upper() if code else "USD")
        if best is None or amount > best[0]:    # prefer the largest (annual base) figure
            best = (round(amount, 2), ccy or "USD")
    return best


def _guess_labeled_date(text: str, labels: tuple[str, ...]) -> str | None:
    """A date following one of `labels` (e.g. 'start date'), normalized to ISO, or None."""
    for label in labels:
        m = re.search(re.escape(label) + r"\s*[:\-]?\s*" + _DATE_TOKEN, text, re.IGNORECASE)
        if m:
            try:
                return normalize.to_iso_date(m.group(1))
            except ValueError:
                continue
    return None


def _guess_employment_type(text: str) -> str | None:
    low = text.lower()
    if any(w in low for w in ("part-time", "part time")):
        return "part_time"
    if any(w in low for w in ("contractor", "fixed-term", "fixed term", "temporary", "consultant", "contingent")):
        return "contractor"
    if any(w in low for w in ("full-time", "full time", "permanent", "cdi", "indefinite")):
        return "full_time"
    return None


def _guess_rating(text: str) -> float | None:
    m = re.search(r"(?:overall\s+|performance\s+)?rating\s*[:\-]?\s*([0-5](?:\.[0-9])?)\s*(?:/\s*5)?",
                  text, re.IGNORECASE)
    if m:
        try:
            v = float(m.group(1))
            return v if 0 < v <= 5 else None
        except ValueError:
            return None
    return None


def _guess_goal_attainment(text: str) -> float | None:
    """Goal attainment as a 0–1 fraction (e.g. '95%' -> 0.95)."""
    m = re.search(r"goal\s+attainment\s*[:\-]?\s*([0-9]{1,3})\s*%?", text, re.IGNORECASE)
    if m:
        try:
            v = float(m.group(1))
            return round(v / 100, 2) if v > 1 else round(v, 2)
        except ValueError:
            return None
    return None


def _guess_prior_employers(text: str) -> int | None:
    """Loose proxy for prior roles: count year-range lines ('2018 – 2021', '2019–present')."""
    n = len(_DATE_RANGE_RE.findall(text))
    return n if n > 0 else None


def _section_list(lines: list[str], header_re: str, cap: int) -> list[str]:
    """Read a comma/semicolon list from a 'Header: a, b, c' line or a standalone header's
    following lines (bounded by the next section header)."""
    for i, line in enumerate(lines):
        m = re.match(header_re + r"\s*[:\-]?\s*(.*)$", line, re.IGNORECASE)
        if not m:
            continue
        inline = m.group(1).strip()
        if inline:
            blob = inline
        else:
            collected: list[str] = []
            for nxt in lines[i + 1:i + 5]:
                first = nxt.split()[0].lower().strip(":,") if nxt.split() else ""
                if first in _SECTION_WORDS:
                    break
                collected.append(nxt)
            blob = " , ".join(collected)
        parts = re.split(r"[,;|/•·]+", blob)
        items = _dedupe_keep_order([p.strip() for p in parts if 1 < len(p.strip()) <= 40])
        if items:
            return items[:cap]
    return []


def _guess_languages(text: str, lines: list[str]) -> list[str]:
    found = _section_list(lines, r"\s*languages?", 8)
    if found:
        return found
    low = text.lower()
    hits = [lng.title() for lng in _LANGUAGES if re.search(rf"(?<!\w){lng}(?!\w)", low)]
    return _dedupe_keep_order(hits)[:8]


def _guess_certifications(text: str, lines: list[str]) -> list[str]:
    return _section_list(lines, r"\s*certifications?", 10)


def _guess_university(text: str) -> str | None:
    for rx in (_UNI_RE, _UNI_OF_RE):
        m = rx.search(text)
        if m:
            return m.group(1).strip()
    return None


def _guess_education_field(text: str) -> str | None:
    low = text.lower()
    for field, kws in _EDU_FIELD_KEYWORDS:
        if any(k in low for k in kws):
            return field
    return None


def suggest_fields(text: str, doc_type: str = "auto", filename: str = "") -> dict[str, Any]:
    """Propose employee fields from extracted document text, tailored to the document type
    (auto-detected when `doc_type` is "auto"). Returns {field: {value, confidence}} for the
    fields found, plus skills/experience/education context and the resolved `doc_type`. Missing
    fields are omitted — never invented.
    """
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    dtype = detect_doc_type(filename, text) if doc_type in (None, "", "auto") else doc_type
    suggestions: dict[str, dict[str, Any]] = {}

    def put(field: str, value: Any, confidence: float) -> None:
        if value not in (None, "", []):
            suggestions[field] = {"value": value, "confidence": round(confidence, 2)}

    # Identity + contact: every document tends to name the person, so always attempt these.
    email_m = _EMAIL_RE.search(text or "")
    if email_m:
        put("email", email_m.group(0).lower(), 0.96)
    phone_m = _PHONE_RE.search(text or "")
    if phone_m:
        put("phone", phone_m.group(1).strip(), 0.8)
    put("full_name", _guess_name(lines), 0.6)
    role = _guess_role(lines)
    if role:
        put("role", role, 0.55)
        put("title", role, 0.55)
    put("location", _guess_location(lines), 0.6)

    skills: list[str] = []
    languages: list[str] = []
    certifications: list[str] = []
    years = education = None

    if dtype == "cv":
        skills = _guess_skills(text or "", lines)
        languages = _guess_languages(text or "", lines)
        certifications = _guess_certifications(text or "", lines)
        years = _guess_years(text or "")
        education = _guess_education(text or "")
        put("education_field", _guess_education_field(text or ""), 0.5)
        put("education_institution", _guess_university(text or ""), 0.55)
        put("years_of_experience", years, 0.6)
        if years:
            put("total_working_years", years, 0.5)   # a sensible seed the operator can adjust
        put("prior_employer_count", _guess_prior_employers(text or ""), 0.45)

    if dtype in ("contract", "offer", "payslip"):
        money = _guess_money(text or "")
        if money:
            put("base_salary", money[0], 0.7 if dtype != "payslip" else 0.6)
            put("currency", money[1], 0.85)
        if dtype != "payslip":
            put("hire_date", _guess_labeled_date(text or "", _HIRE_LABELS), 0.7)
            put("employment_type", _guess_employment_type(text or ""), 0.7)

    if dtype == "review":
        r = _guess_rating(text or "")
        if r is not None:
            put("performance_rating", r, 0.7)
        g = _guess_goal_attainment(text or "")
        if g is not None:
            put("goal_attainment", g, 0.7)

    return {
        "fields": suggestions,
        "skills": skills,
        "languages": languages,
        "certifications": certifications,
        "years_experience": years,
        "education_level": education,
        "doc_type": dtype,
        "excerpt": "\n".join(lines[:6]),
        "char_count": len(text or ""),
    }
